from fasthtml.common import *
import os
import requests
import json
import uuid
import asyncio
import time
import re
from datetime import datetime
from dotenv import load_dotenv
from auth import validate_password_strength, hash_password, verify_password, generate_totp_secret, verify_totp

load_dotenv()

# --- THEME & SCRIPTS ---
theme_hdrs = [
    Script(src="https://cdn.tailwindcss.com"),
    Script(src="https://cdn.jsdelivr.net/npm/marked/marked.min.js"),
    Script(src="https://cdnjs.cloudflare.com/ajax/libs/highlight.js/11.9.0/highlight.min.js"),
    Link(rel="stylesheet", href="https://cdnjs.cloudflare.com/ajax/libs/highlight.js/11.9.0/styles/tokyo-night-dark.min.css"),
    Style("""
        body { background-color: #0f172a; color: #f8fafc; font-family: system-ui, sans-serif; margin: 0;}
        .sidebar { background-color: #1e293b; border-right: 1px solid #333; height: 100vh; padding: 20px; display: flex; flex-direction: column;}
        .main-content { padding: 20px; height: 100vh; display: flex; flex-direction: column; position: relative;}
        .chat-container { flex-grow: 1; overflow-y: auto; margin-bottom: 20px; border: 1px solid #333; padding: 15px; border-radius: 8px; background: #151e2e; }
        .input-row { display: flex; gap: 10px; align-items: flex-end; width: 100%; background: #1e293b; padding: 10px; border-radius: 10px; border: 1px solid #333;}
        .input-box { flex-grow: 1; padding: 12px; border-radius: 5px; background: #2a2a35; border: 1px solid #444; color: white; outline: none; }
        .file-upload-btn { background: #2a2a35; border: 1px solid #444; color: #94a3b8; padding: 12px; border-radius: 5px; cursor: pointer; font-size:14px; display: flex; align-items:center; justify-content:center;}
        .model-selector { padding: 8px; border-radius: 5px; background: #1e293b; border: 1px solid #444; color: #00e5ff; width: 100%; font-size: 12px; outline:none; }
        .send-btn { padding: 12px 24px; border-radius: 5px; background: #00e5ff; color: #0f172a; font-weight: bold; cursor: pointer; border: none; }
        .msg-user { color: #00e5ff; font-size: 12px; font-weight: bold; margin-bottom: 2px; margin-top: 15px;}
        .msg-ia { color: #a855f7; font-size: 12px; font-weight: bold; margin-bottom: 2px; margin-top: 15px;}
        .bubble-user { background: #1e293b; padding: 10px; border-radius: 5px; display: inline-block; color: white; border: 1px solid #333;}
        .bubble-ia { background: #1e293b; padding: 15px; border-radius: 8px; display: block; color: #e2e8f0; border: 1px solid #333;}
        .bubble-ia pre { background: #111827; padding: 15px; border-radius: 8px; margin-top: 10px; position: relative; border: 1px solid #444;}
        .htmx-indicator { display: none; position: absolute; bottom: 180px; left: 50%; transform: translateX(-50%); background: rgba(15, 23, 42, 0.9); border: 1px solid #00e5ff; padding: 15px 30px; border-radius: 50px; color: #00e5ff; font-weight: bold; z-index: 50;}
        .htmx-request .htmx-indicator { display: block; }
        .history-btn { background: #2a2a35; color: #e2e8f0; border: 1px solid #444; padding: 8px; border-radius: 5px; cursor: pointer; text-align: left; }
        .history-btn:hover { background: #38bdf8; color: #0f172a; }
        .export-btn { display: block; background: #0f172a; border: 1px solid #00e5ff; color: #00e5ff; font-weight: bold; padding: 10px; border-radius: 5px; text-align: center; cursor: pointer; transition: 0.2s; width: 100%;}
        .export-btn:hover { background: #00e5ff; color: #0f172a; }
    """),
    Script("""
        function filterModels() {
            let filter = document.querySelector('input[name="filter_type"]:checked').value;
            document.querySelectorAll('.model-selector option').forEach(opt => {
                if(opt.value === "") return;
                let isFree = opt.getAttribute('data-free') === 'true';
                if(filter === 'free' && !isFree) opt.style.display = 'none';
                else if(filter === 'paid' && isFree) opt.style.display = 'none';
                else opt.style.display = 'block';
            });
        }
        function validateMultiIA() {
            let wCount = 0;
            for(let i=1; i<=5; i++) { if(document.getElementById('worker'+i).value !== "") wCount++; }
            if(wCount === 0) { 
                alert("Erreur : Vous devez sélectionner au moins une IA de Travail."); 
                return false; 
            }
            if(wCount > 1) {
                if(document.getElementById('writer').value === "" || document.getElementById('synthesizer').value === "") {
                    alert("Erreur : Puisque vous avez sélectionné plusieurs travailleurs, le Rédacteur et le Concaténeur sont obligatoires.");
                    return false;
                }
            }
            return true;
        }
        function updateFileCount(input) {
            let label = document.getElementById('file-label-text');
            if(input.files && input.files.length > 0) { label.innerText = `📎 ${input.files.length} fichier(s)`; label.style.color = '#00e5ff'; } 
            else { label.innerText = '📎 Fichiers Joints'; label.style.color = '#94a3b8'; }
        }
        htmx.onLoad(function(content) {
            content.querySelectorAll('.bubble-ia:not(.rendered)').forEach(function(el) {
                el.innerHTML = marked.parse(el.textContent);
                el.classList.add('rendered');
                el.querySelectorAll('pre code').forEach((block) => { hljs.highlightElement(block); });
                let chat = document.getElementById('chat-history'); chat.scrollTop = chat.scrollHeight;
            });
        });
    """)
]

app, rt = fast_app(secret_key=os.getenv("SESSION_SECRET", "super-secret-key-fallback"), hdrs=theme_hdrs)

# --- AUTHENTIFICATION ---
@rt('/login')
def get():
    return Title("Connexion - AETHAS38"), Body(
        Div(
            H2("AETHAS 38 - Accès Sécurisé", style="color:#00e5ff; margin-bottom:20px; text-align:center;"),
            Form(
                Input(type="email", name="email", placeholder="E-mail", cls="input-box mb-3 w-full", required=True),
                Input(type="password", name="password", placeholder="Mot de passe", cls="input-box mb-3 w-full", required=True),
                Input(type="text", name="totp_code", placeholder="Code 2FA (6 chiffres)", cls="input-box mb-4 w-full", required=True),
                Button("Se connecter", type="submit", cls="send-btn w-full"),
                hx_post="/auth/login", hx_target="#auth-response"
            ),
            Div(id="auth-response", style="color:red; margin-top:10px; text-align:center;"),
            cls="sidebar", style="margin: auto; width: 400px; height: auto; border-radius: 10px; margin-top: 15vh;"
        ), style="background-color: #0f172a; height: 100vh; display: flex;"
    )

@rt('/auth/login')
def post(email: str, password: str, totp_code: str, session):
    admin_email = os.getenv("ADMIN_EMAIL")
    admin_hash = os.getenv("ADMIN_PASSWORD_HASH")
    totp_secret = os.getenv("ADMIN_2FA_SECRET")

    if not admin_email or not admin_hash or not totp_secret:
        return Div("Configuration système incomplète (.env).", style="color:red; font-weight:bold;")

    if email == admin_email and verify_password(admin_hash, password) and verify_totp(totp_secret, totp_code):
        session["authenticated"] = True
        session["role"] = "admin"
        return Script("window.location.href = '/';")
    else:
        return Div("Identifiants ou code 2FA incorrects.", style="color:#ef4444; font-weight:bold;")

@rt('/logout')
def get(session):
    session.clear()
    return RedirectResponse('/login')

# --- SYSTEM & LOGS ---
@rt('/check_updates')
def get(session):
    if not session.get("authenticated"): return Span("")
    try:
        res = requests.get("https://gitea.aethas38.duckdns.org/api/v1/repos/xavier/MULTI-IA-CODAGE/commits?limit=1", timeout=3)
        if res.status_code == 200:
            remote_commit = res.json()[0].get("sha", "")[:7]
            return Span(f"🟢 Gitea Connecté (Dernier commit : {remote_commit})", style="color:#10b981;")
        return Span("🟠 Gitea injoignable", style="color:#fbbf24;")
    except Exception as e:
        return Span("⚪ Statut réseau inconnu", style="color:#94a3b8;")

# Création des dossiers qui seront mappés sur le RAID
os.makedirs("sessions", exist_ok=True)
os.makedirs("logs", exist_ok=True)
APP_LAUNCH_TIME = datetime.now().strftime("%Y%m%d_%H%M%S")
LOG_FILE = f"logs/AETHAS38_Run_{APP_LAUNCH_TIME}.log"

def log_event(session_id, category, action):
    sid_short = session_id[:8] if session_id and session_id != "system" else "SYSTEM"
    timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    with open(LOG_FILE, "a", encoding="utf-8") as f:
        f.write(f"[{timestamp}] [Session:{sid_short}] [{category}] {action}\n")

# --- CATALOGUE MANAGER (Sécurisé sur le RAID) ---
CATALOG_FILE = "sessions/catalog.json"
MODELS_DATA = []

def load_catalog():
    global MODELS_DATA
    if os.path.exists(CATALOG_FILE):
        with open(CATALOG_FILE, "r", encoding="utf-8") as f:
            try: MODELS_DATA = json.load(f)
            except: MODELS_DATA = []
    if not MODELS_DATA:
        # Fallback de sécurité mis à jour
        MODELS_DATA = [
            {"id": "groq|llama-3.1-8b-instant", "name": "Groq - Llama 3.1 (8B)", "is_free": True, "provider": "Général / Texte", "description": "Modèle par défaut."},
            {"id": "gemini|gemini-1.5-flash", "name": "Google - Gemini 1.5 Flash", "is_free": True, "provider": "Vision", "description": "Modèle par défaut."}
        ]
load_catalog()

def save_catalog(data):
    global MODELS_DATA
    MODELS_DATA = data
    with open(CATALOG_FILE, "w", encoding="utf-8") as f:
        json.dump(data, f, indent=4, ensure_ascii=False)

def get_budget():
    mgmt_key = os.getenv("OPENROUTER_MANAGEMENT_KEY")
    if not mgmt_key: return "Clé manquante"
    try:
        data = requests.get("https://openrouter.ai/api/v1/credits", headers={"Authorization": f"Bearer {mgmt_key}"}, timeout=5).json().get("data", {})
        return f"{data.get('total_credits', 0) - data.get('total_usage', 0):.4f} $"
    except: return "Erreur réseau"

# --- PERSISTENCE HISTORIQUE (Sécurisé sur le RAID) ---
HISTORY_FILE = "sessions/history.json"

def load_index():
    if not os.path.exists(HISTORY_FILE): return {}
    with open(HISTORY_FILE, "r", encoding="utf-8") as f:
        try: return json.load(f)
        except: return {}
def save_index(idx):
    with open(HISTORY_FILE, "w", encoding="utf-8") as f:
        json.dump(idx, f, indent=4, ensure_ascii=False)
def load_session_data(session_id):
    path = f"sessions/{session_id}.json"
    if not os.path.exists(path): return []
    with open(path, "r", encoding="utf-8") as f:
        try: return json.load(f)
        except: return []
def save_session_data(session_id, messages):
    with open(f"sessions/{session_id}.json", "w", encoding="utf-8") as f:
        json.dump(messages, f, indent=4, ensure_ascii=False)
def create_new_session():
    session_id = str(uuid.uuid4())
    idx = load_index()
    idx[session_id] = {"title": f"Projet du {datetime.now().strftime('%d/%m %H:%M')}", "pinned": False}
    save_index(idx)
    save_session_data(session_id, [])
    return session_id

def render_history_list():
    idx = load_index()
    sorted_sessions = sorted(idx.items(), key=lambda x: (not x[1].get('pinned', False), x[1]['title']), reverse=False)
    items = [Button("➕ Nouveau Projet", cls="export-btn mb-3", style="background:#10b981; color:#0f172a;", hx_get="/session/new", hx_target="#chat-history")]
    for sid, data in sorted_sessions:
        pin_icon, pin_color = ("📍", "#10b981") if data.get('pinned') else ("📌", "#94a3b8")
        item = Div(
            Div(data['title'], cls="truncate flex-grow cursor-pointer text-sm hover:text-cyan-400", hx_get=f"/session/load/{sid}", hx_target="#chat-history"),
            Div(
                Button("✏", cls="text-xs mx-1 hover:text-white", onclick=f"let name = prompt('Nouveau nom:'); if(name) {{ htmx.ajax('POST', '/history/rename/{sid}', {{values: {{title: name}}, target: '#history-list-container'}}); }}"),
                Button(pin_icon, style=f"color:{pin_color};", cls="text-xs mx-1 hover:text-white", hx_post=f"/history/pin/{sid}", hx_target="#history-list-container"),
                Button("🗑️", cls="text-xs hover:text-red-500", onclick=f"if(confirm('Supprimer ce projet ?')) {{ htmx.ajax('POST', '/history/delete/{sid}', {{target: '#history-list-container'}}); }}"),
                cls="flex-shrink-0"
            ),
            cls="flex justify-between items-center w-full mb-2 bg-[#2a2a35] p-2 rounded border border-[#444]"
        )
        items.append(item)
    return Div(*items, id="history-list-container", hx_swap_oob="true")

# --- UI COMPONENTS ---
def render_model_dropdown(field_id, placeholder):
    groups = {}
    for m in MODELS_DATA:
        prov = m.get('provider', 'Général')
        if prov not in groups: groups[prov] = []
        groups[prov].append(m)
    
    opts = [Option("--- Laisser vide ---", value="")]
    for prov in sorted(groups.keys()):
        grp_opts = []
        for m in sorted(groups[prov], key=lambda x: x['name']):
            is_free = str(m.get('is_free', False)).lower()
            label = f"{m['name']} ({'Gratuit' if is_free=='true' else 'Payant'})"
            grp_opts.append(Option(label, value=m['id'], **{"data-free": is_free}))
        opts.append(Optgroup(label=prov)(*grp_opts))
        
    return Div(
        P(placeholder, style="color:#94a3b8; font-size:10px; font-weight:bold; margin-bottom:2px;"),
        Select(*opts, name=field_id, id=field_id, cls="model-selector", onchange="filterModels()"),
        style="margin-bottom: 8px;"
    )

def parse_txt_catalog(content: str):
    """ Logique VBA traduite en Python pour le fallback du fichier TXT """
    lines = content.split('\n')
    models = []
    current = {}
    ext_date = "Inconnue"
    ext_time = ""
    
    for line in lines:
        line_trim = line.strip()
        
        if "Généré le" in line_trim:
            match = re.search(r"Généré le (\d{2}/\d{2}/\d{4}) à (\d{2}:\d{2}:\d{2})", line_trim)
            if match:
                ext_date = match.group(1)
                ext_time = match.group(2)
        
        if line_trim.startswith("Nom : "): current['name'] = line_trim.replace("Nom : ", "").strip()
        elif line_trim.startswith("ID : "): current['id'] = line_trim.replace("ID : ", "").strip()
        elif line_trim.startswith("Tarif : "): current['is_free'] = ("Gratuit" in line_trim)
        elif line_trim.startswith("Spécialité : "): current['provider'] = line_trim.replace("Spécialité : ", "").strip()
        elif line_trim.startswith("Description : "): current['description'] = line_trim.replace("Description : ", "").strip()
        elif line_trim.startswith("---"):
            if 'id' in current and 'name' in current: models.append(current)
            current = {}
            
    return models, ext_date, ext_time

# --- API DISTANTE ---
def ask_llm(session_id, full_id, msg):
    parts = full_id.split("|", 1)
    provider = parts[0]
    actual_model = parts[1] if len(parts) > 1 else full_id
    
    log_event(session_id, "API_REQ", f"Provider: {provider} | Modèle: {actual_model}")
    try:
        if provider == "openrouter":
            res = requests.post("https://openrouter.ai/api/v1/chat/completions", headers={"Authorization": f"Bearer {os.getenv('OPENROUTER_API_KEY')}"}, json={"model": actual_model, "messages": [{"role": "user", "content": msg}]})
            if res.status_code == 200: return res.json()["choices"][0]["message"]["content"]
            return f"⚠️ Erreur API ({provider}) : {res.text}"
        elif provider == "groq":
            res = requests.post("https://api.groq.com/openai/v1/chat/completions", headers={"Authorization": f"Bearer {os.getenv('GROQ_API_KEY')}"}, json={"model": actual_model, "messages": [{"role": "user", "content": msg}]})
            if res.status_code == 200: return res.json()["choices"][0]["message"]["content"]
            return f"⚠️ Erreur API ({provider}) : {res.text}"
        elif provider == "gemini":
            res = requests.post(f"https://generativelanguage.googleapis.com/v1beta/models/{actual_model}:generateContent?key={os.getenv('GEMINI_API_KEY')}", json={"contents": [{"parts": [{"text": msg}]}]})
            if res.status_code == 200: return res.json()["candidates"][0]["content"]["parts"][0]["text"]
            return f"⚠️ Erreur API ({provider}) : {res.text}"
        return "⚠️ Fournisseur API non implémenté ou ID invalide."
    except Exception as e: return f"⚠️️ Erreur de connexion : {str(e)}"

async def async_ask_llm(session_id, full_id, msg):
    return await asyncio.to_thread(ask_llm, session_id, full_id, msg)

# --- ROUTES & VIEWS ---
SESSION_PROGRESS = {}

@rt('/status/{sid}')
def get_status(sid: str):
    return SESSION_PROGRESS.get(sid, "⚙️ Traitement en cours...")

@rt('/')
def get(session):
    if not session.get("authenticated"): return RedirectResponse('/login')
    session_id = create_new_session()
    return Title("AETHAS38 Orchestrateur"), Body(
        Div(
            # --- SIDEBAR GAUCHE ---
            Div(
                H2("AETHAS 38", style="color:#00e5ff; font-weight:bold; font-size:24px;"),
                P("Orchestrateur Distant", style="color:#a855f7; margin-bottom: 10px; font-size:12px;"),
                Div("⏳ Vérification GitHub...", hx_get="/check_updates", hx_trigger="load", style="font-size:10px; margin-bottom:15px;"),
                
                Div(
                    H3("Importer Nouveaux Modèles", style="color:#94a3b8; font-size:12px; font-weight:bold; margin-bottom:5px;"),
                    Form(
                        Input(type="file", name="catalog_file", accept=".txt,.xlsm,.xlsx", required=True, id="cat-upload", style="display:none;", onchange="document.getElementById('cat-label').innerText = '📎 Fichier Sélectionné';"),
                        Label(Span("📎 Sélectionner (.txt ou .xlsm)", id="cat-label"), _for="cat-upload", cls="file-upload-btn mb-2"),
                        Button("Importer le Catalogue", type="submit", cls="export-btn", style="background:#a855f7; color:#fff; border:none;"),
                        hx_post="/import_catalog", hx_target="#catalog-status", enctype="multipart/form-data"
                    ),
                    Div(id="catalog-status", style="font-size:10px; color:#10b981; margin-top:5px; text-align:center;"),
                    style="background:#151e2e; padding:10px; border-radius:5px; border:1px solid #333; margin-bottom:15px;"
                ),

                Div(H3("Historique", style="color:#94a3b8; font-size:12px; font-weight:bold; margin-bottom:10px;"), render_history_list(), style="flex-grow:1; overflow:hidden;"),
                
                Button("📥 Exporter Modèles Actuels (.txt)", type="button", cls="export-btn", onclick="window.location.href='/export_models'"),
                A("Déconnexion", href="/logout", cls="export-btn w-full mt-2", style="color:#ef4444; border-color:#ef4444; margin-bottom:10px;"),
                Div(P("BUDGET", style="color:#94a3b8; font-size:10px; font-weight:bold;"), P(get_budget(), id="budget-display", style="color:#10b981; font-size:16px; font-weight:bold;")),
                cls="sidebar"
            ),
            # --- ZONE PRINCIPALE ---
            Div(
                Div(id="chat-history", cls="chat-container"),
                Div(Span("⚙️ Préparation du Pipeline...", id="loading-text"), id="loading-tracker", cls="htmx-indicator"),
                
                Form(
                    Input(type="hidden", name="session_id", value=session_id, id="current-session-id"),
                    
                    Div(
                        P("Filtre d'affichage :", style="color:#00e5ff; font-weight:bold; font-size:12px; margin-right:15px;"),
                        Label(Input(type="radio", name="filter_type", value="all", checked=True, onchange="filterModels()"), " Tous"),
                        Label(Input(type="radio", name="filter_type", value="free", onchange="filterModels()"), " Gratuits Uniquement"),
                        Label(Input(type="radio", name="filter_type", value="paid", onchange="filterModels()"), " Payants Uniquement"),
                        cls="flex items-center text-sm text-white gap-4 mb-4 bg-[#1e293b] p-2 rounded border border-[#333] w-fit"
                    ),
                    
                    Div(
                        Div(render_model_dropdown("writer", "1. Rédacteur de Prompt"), render_model_dropdown("synthesizer", "7. Concaténeur Final"), style="border-right: 1px solid #444; padding-right:15px;"),
                        Div(
                            render_model_dropdown("worker1", "2. IA Travailleur 1 (Obligatoire)"),
                            render_model_dropdown("worker2", "3. IA Travailleur 2"),
                            render_model_dropdown("worker3", "4. IA Travailleur 3"),
                            style="padding: 0 15px;"
                        ),
                        Div(
                            render_model_dropdown("worker4", "5. IA Travailleur 4"),
                            render_model_dropdown("worker5", "6. IA Travailleur 5"),
                            style="padding-left: 15px;"
                        ),
                        cls="grid grid-cols-3 gap-4 mb-4 bg-[#151e2e] p-4 rounded border border-[#333]"
                    ),
                    
                    Div(
                        Input(type="file", name="fichiers", id="file-upload", multiple=True, style="display:none;", onchange="updateFileCount(this)"),
                        Label(Span("📎 Fichiers Joints", id="file-label-text"), _for="file-upload", cls="file-upload-btn w-40"),
                        Input(type="text", name="msg", placeholder="Détaillez le travail à effectuer...", cls="input-box", required=True),
                        Button("Lancer le Pipeline", type="submit", cls="send-btn"),
                        cls="input-row"
                    ),
                    hx_post="/chat", hx_target="#chat-history", hx_swap="beforeend", hx_indicator="#loading-tracker", enctype="multipart/form-data",
                    onsubmit="return validateMultiIA();",
                    **{"hx-on:htmx:afterRequest": "if(event.detail.successful) { this.reset(); document.getElementById('file-label-text').innerText='📎 Fichiers Joints'; filterModels(); }"}
                ),
                cls="main-content"
            ),
            style="display: grid; grid-template-columns: 280px 1fr;"
        ),
        Script("""
            let statusInterval;
            document.addEventListener('htmx:beforeRequest', function(evt) {
                if(evt.detail.target.id === 'chat-history') {
                    let sid = document.getElementById('current-session-id').value;
                    document.getElementById('loading-text').innerText = "⚙️ Envoi des données...";
                    statusInterval = setInterval(async () => {
                        try { let res = await fetch('/status/' + sid); if(res.ok) document.getElementById('loading-text').innerText = await res.text(); } catch(e) {}
                    }, 800);
                }
            });
            document.addEventListener('htmx:afterRequest', function(evt) { if(evt.detail.target.id === 'chat-history') clearInterval(statusInterval); });
            filterModels();
        """)
    )

# --- PIPELINE DE TRAITEMENT ---
@rt('/chat')
async def post(msg: str, session_id: str, session, writer: str = "", worker1: str = "", worker2: str = "", worker3: str = "", worker4: str = "", worker5: str = "", synthesizer: str = "", fichiers: list[UploadFile] = None):
    if not session.get("authenticated"): return RedirectResponse('/login')
    t_start = time.time()
    
    files_context, noms_fichiers = "", []
    if fichiers:
        for f in fichiers:
            if f.filename:
                noms_fichiers.append(f.filename)
                try: files_context += f"\n--- {f.filename} ---\n{(await f.read()).decode('utf-8')}\n"
                except: files_context += f"\n--- {f.filename} (Binaire ignoré) ---\n"
    
    base_prompt = f"Fichiers fournis :\n{files_context}\n\nConsigne de l'utilisateur : {msg}" if files_context else msg
    workers = [w for w in [worker1, worker2, worker3, worker4, worker5] if w]
    
    t_w0 = time.time()
    if writer:
        SESSION_PROGRESS[session_id] = "🧠 Étape 1 : Le Rédacteur prépare l'invite..."
        prompt_to_work = ask_llm(session_id, writer, "Tu es un Ingénieur Prompt Senior. Reformule et optimise techniquement cette demande pour des LLM spécialisés. Retourne uniquement l'invite optimisée sans salutations :\n\n" + base_prompt)
    else:
        prompt_to_work = base_prompt
    t_w1 = time.time()

    SESSION_PROGRESS[session_id] = f"⚡ Étape 2 : Travail en cours ({len(workers)} IA en parallèle)..."
    tasks = [async_ask_llm(session_id, w, prompt_to_work) for w in workers]
    results = await asyncio.gather(*tasks)
    t_w2 = time.time()

    if synthesizer and len(workers) > 0:
        SESSION_PROGRESS[session_id] = "🏗️ Étape 3 : Synthèse et assemblage final..."
        synth_input = "Tu es l'Architecte Final. Voici la même tâche effectuée par plusieurs intelligences artificielles :\n\n"
        for i, res in enumerate(results): synth_input += f"--- Proposition IA {i+1} ---\n{res}\n\n"
        synth_input += "Analyse ces propositions, garde le meilleur code, corrige les erreurs potentielles et génère la solution finale absolue et complète en français."
        final_response = ask_llm(session_id, synthesizer, synth_input)
    else:
        final_response = results[0]
    t_w3 = time.time()
    SESSION_PROGRESS[session_id] = "✅ Pipeline Terminé"

    rapport = (
        f"\n\n---\n⏱️ **Rapport de Performance (Pipeline Distant)**\n"
        f"- **Rédacteur** : {t_w1 - t_w0:.2f} s\n"
        f"- **Travailleurs (Parallèle)** : {t_w2 - t_w1:.2f} s\n"
        f"- **Concaténeur** : {t_w3 - t_w2:.2f} s\n"
        f"- **Temps Total du Cycle** : {t_w3 - t_start:.2f} s"
    )

    info_lbl = f"<br><span style='color:#38bdf8; font-size:10px;'>⚙️ {len(workers)} Travailleur(s) | 📎 {len(noms_fichiers)} fichier(s)</span>"
    
    session_data = load_session_data(session_id)
    session_data.append({"user": msg, "ia": final_response + rapport, "info": info_lbl})
    save_session_data(session_id, session_data)

    return Div(P("Vous", cls="msg-user"), Div(NotStr(f"{msg}{info_lbl}"), cls="bubble-user"), P("AETHAS38", cls="msg-ia"), Div(final_response + rapport, cls="bubble-ia")), P(get_budget(), id="budget-display", style="color:#10b981; font-size:16px; font-weight:bold;", hx_swap_oob="true")

# --- IMPORT / EXPORT CATALOGUE ---
# --- IMPORT / EXPORT CATALOGUE ---
@rt('/export_models')
def get(session):
    if not session.get("authenticated"): return RedirectResponse('/login')
    log_event("system", "SYSTEM", "Demande d'exportation du catalogue depuis OpenRouter.")
    now = datetime.now()
    txt_content = f"=== CATALOGUE COMPLET DES MODÈLES (Généré le {now.strftime('%d/%m/%Y à %H:%M:%S')}) ===\n\n"
    
    # 1. Ajout des modèles de base gratuits
    models_to_export = [
        {"id": "groq|llama-3.1-8b-instant", "name": "Groq - Llama 3.1 (8B)", "pricing": {"prompt": 0, "completion": 0}, "architecture": {"modality": "text"}, "description": "Modèle ultra-rapide hébergé par Groq."},
        {"id": "gemini|gemini-1.5-flash", "name": "Google - Gemini 1.5 Flash", "pricing": {"prompt": 0, "completion": 0}, "architecture": {"modality": "text/vision"}, "description": "Modèle multimodal léger de Google."}
    ]
    
    # 2. Téléchargement en direct de tous les modèles d'OpenRouter
    try:
        res = requests.get("https://openrouter.ai/api/v1/models", timeout=10)
        if res.status_code == 200:
            for m in res.json().get("data", []):
                m["id"] = f"openrouter|{m['id']}"
                m["name"] = f"OR - {m['name']}"
                models_to_export.append(m)
    except Exception as e:
        log_event("system", "ERROR", f"Erreur API OpenRouter lors de l'export : {e}")

    google_banned = False
    
    for m in models_to_export:
        name = m.get('name', 'Inconnu')
        m_id = m.get('id', 'Inconnu')
        
        usage = "Général / Texte"
        arch_mod = str(m.get('architecture', {}).get('modality', '')).lower()
        m_id_lower = m_id.lower()
        
        if "vision" in arch_mod or "vision" in m_id_lower or "vl" in m_id_lower or "image" in arch_mod:
            usage = "Vision (Analyse d'images & Multimodal)"
        elif "coder" in m_id_lower or "code" in m_id_lower or "math" in m_id_lower:
            usage = "Codage & Mathématiques"
        elif "video" in m_id_lower:
            usage = "Vidéo"
        elif "rp" in m_id_lower or "uncensored" in m_id_lower or "roleplay" in m_id_lower:
            usage = "Non-censuré / Roleplay"
            
        is_free = m.get("is_free")
        if is_free is None:
            pricing = m.get('pricing', {})
            p_prompt = float(pricing.get('prompt', -1))
            p_comp = float(pricing.get('completion', -1))
            cout = "Gratuit" if p_prompt == 0.0 and p_comp == 0.0 else f"Payant (In: {p_prompt}, Out: {p_comp})"
        else:
            cout = "Gratuit" if is_free else "Payant"
            
        desc_courte = str(m.get('description', '')).replace('\n', ' ')[:150]
        desc_fr = desc_courte
        
        if "openrouter" in m_id and desc_courte:
            if not google_banned:
                try:
                    url = "https://translate.googleapis.com/translate_a/single?client=gtx&sl=en&tl=fr&dt=t&q=" + requests.utils.quote(desc_courte)
                    trans_res = requests.get(url, timeout=1.5)
                    if trans_res.status_code == 200:
                        desc_fr = "".join([x[0] for x in trans_res.json()[0]])
                    else:
                        google_banned = True
                        desc_fr += " [Traduction interrompue: Sécurité anti-spam]"
                except: 
                    google_banned = True
                    desc_fr += " [Traduction interrompue: Délai d'attente]"
            else:
                desc_fr = desc_courte
                
        if not desc_fr: desc_fr = "Aucune description fournie."
        
        txt_content += (
            f"Nom : {name}\n"
            f"ID : {m_id}\n"
            f"Tarif : {cout}\n"
            f"Spécialité : {usage}\n"
            f"Description : {desc_fr}\n"
            f"{'-'*60}\n"
        )

    filename = f"modeles_ia_{datetime.now().strftime('%Y%m%d_%H%M')}.txt"
    return Response(txt_content, media_type="text/plain", headers={"Content-Disposition": f"attachment; filename={filename}"})

@rt('/import_catalog')
async def post(catalog_file: UploadFile, session):
    if not session.get("authenticated"): return RedirectResponse('/login')
    filename = catalog_file.filename.lower()
    new_models = []
    ext_date, ext_time = "Inconnue", ""
    
    if filename.endswith(".txt"):
        content = (await catalog_file.read()).decode('utf-8').replace('\r\n', '\n')
        new_models, ext_date, ext_time = parse_txt_catalog(content)
        
    elif filename.endswith((".xlsx", ".xlsm")):
        try:
            import openpyxl
            from io import BytesIO
            content = await catalog_file.read()
            wb = openpyxl.load_workbook(BytesIO(content), data_only=True)
            ws = wb.active
            
            val_b2 = ws["B2"].value
            val_c2 = ws["C2"].value
            
            if val_b2 and hasattr(val_b2, 'strftime'): ext_date = val_b2.strftime('%d/%m/%Y')
            else: ext_date = str(val_b2) if val_b2 else "Inconnue"
            
            if val_c2 and hasattr(val_c2, 'strftime'): ext_time = val_c2.strftime('%H:%M:%S')
            else: ext_time = str(val_c2) if val_c2 else ""
            
            for row in ws.iter_rows(min_row=5, values_only=True):
                if row[0] and row[1]:
                    new_models.append({
                        'name': str(row[0]).strip(),
                        'id': str(row[1]).strip(),
                        'is_free': ("Gratuit" in str(row[2])) if row[2] else False,
                        'provider': str(row[3]).strip() if row[3] else "Général",
                        'description': str(row[4]).strip() if row[4] else ""
                    })
        except ImportError:
            return Div("❌ Erreur : 'openpyxl' n'est pas installé sur le serveur. Veuillez uploader le fichier .txt.", style="color:red;")
        except Exception as e:
            return Div(f"❌ Erreur Excel : {str(e)}", style="color:red;")
    else:
        return Div("❌ Format invalide (.txt ou .xlsm attendu)", style="color:red;")
        
    if new_models:
        save_catalog(new_models)
        date_str = f"Extrait le {ext_date} à {ext_time}." if ext_date != "Inconnue" else ""
        return Div(f"✅ Catalogue mis à jour ({len(new_models)} modèles). {date_str}", Script("setTimeout(()=>window.location.reload(), 2000);"), style="color:#10b981;")
    else:
        return Div("❌ Aucun modèle valide trouvé.", style="color:red;")

# --- HISTORY SYSTEM ---
@rt('/session/new')
def get(session):
    if not session.get("authenticated"): return RedirectResponse('/login')
    sid = create_new_session()
    return Div(), Input(type="hidden", name="session_id", value=sid, id="current-session-id", hx_swap_oob="true"), render_history_list()

@rt('/session/load/{sid}')
def get(sid: str, session):
    if not session.get("authenticated"): return RedirectResponse('/login')
    messages = load_session_data(sid)
    bubbles = [Div(P("Vous", cls="msg-user"), Div(NotStr(f"{msg['user']}{msg.get('info', '')}"), cls="bubble-user"), P("AETHAS38", cls="msg-ia"), Div(msg['ia'], cls="bubble-ia")) for msg in messages]
    return tuple(bubbles) + (Input(type="hidden", name="session_id", value=sid, id="current-session-id", hx_swap_oob="true"),)

@rt('/history/pin/{sid}')
def post_pin(sid: str, session):
    if not session.get("authenticated"): return RedirectResponse('/login')
    idx = load_index()
    if sid in idx:
        idx[sid]['pinned'] = not idx[sid].get('pinned', False)
        save_index(idx)
    return render_history_list()

@rt('/history/rename/{sid}')
def post_rename(sid: str, title: str, session):
    if not session.get("authenticated"): return RedirectResponse('/login')
    idx = load_index()
    if sid in idx:
        idx[sid]['title'] = title
        save_index(idx)
    return render_history_list()

@rt('/history/delete/{sid}')
def post_delete(sid: str, session):
    if not session.get("authenticated"): return RedirectResponse('/login')
    idx = load_index()
    if sid in idx:
        del idx[sid]
        save_index(idx)
        if os.path.exists(f"sessions/{sid}.json"): os.remove(f"sessions/{sid}.json")
    return render_history_list()

if __name__ == '__main__':
    import uvicorn
    uvicorn.run(app, host="0.0.0.0", port=5001)