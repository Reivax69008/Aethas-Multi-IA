from fastapi import FastAPI, Depends, HTTPException, status, Request, Response, UploadFile, File
from fastapi.responses import RedirectResponse, FileResponse, JSONResponse
from fastapi.staticfiles import StaticFiles
from sqlalchemy.orm import Session
from sqlalchemy import text
from typing import List, Dict
import os, json, asyncio, shutil, base64, io, csv, zipfile, traceback, httpx
from datetime import datetime, timedelta, timezone
import pytz

from .database import engine, Base, get_db, SessionLocal
from .auth import get_password_hash, generate_totp_secret, get_totp_uri, verify_password, verify_totp, create_access_token, verify_token
from .schemas import (
    AdminCreate, LoginRequest, ProjectCreate, ProjectResponse, ProjectRename,
    ProjectConfigUpdate, MessageCreate, MessageResponse, PasswordChange,
    ModelReplacementRequest, LogRequest, SystemSettingsUpdate, SystemSettingsResponse,
    MediaGenerateRequest, MediaGenerateResponse
)
from .models import User, Project, Message, SystemSettings, AIModel, FinancialLog
from .orchestrator import (
    run_orchestrator, sync_providers_models, sync_finances, activity_logs,
    log_activity, generate_media, resolve_gemini_model_id, is_worker_eligible,
    find_quota_fallback_model, purge_unauthorized_models, ask_agent_with_fallback
)
from .logger import log_file, cleanup_old_logs, system_logger

Base.metadata.create_all(bind=engine)

try:
    with engine.begin() as conn:
        conn.execute(text("ALTER TABLE projects ADD COLUMN IF NOT EXISTS config_json TEXT;"))
except Exception:
    try:
        with engine.begin() as conn:
            conn.execute(text("ALTER TABLE projects ADD COLUMN config_json TEXT;"))
    except Exception:
        pass

app = FastAPI(title="AETHAS38")

# --- CONFIGURATION HERMES AGENT AUTO-HÉBERGÉ ---
HERMES_API_URL = os.getenv("HERMES_API_URL", "http://192.168.1.20:8642")
HERMES_API_KEY = os.getenv("HERMES_API_KEY", "87186fe63b1e5d8b8d01bc769f5e894d24081a86d0d800d00b6358c26fded37e")

# --- REGISTRE DES TÂCHES ACTIVES POUR INTERRUPTION IMMÉDIATE ---
active_project_tasks: Dict[int, asyncio.Task] = {}

@app.exception_handler(Exception)
async def global_json_exception_handler(request: Request, exc: Exception):
    tb = traceback.format_exc()
    system_logger.error(f"Exception non gérée sur {request.url.path} :\n{tb}")
    log_activity(f"[Erreur Serveur] {str(exc)[:100]}")
    return JSONResponse(
        status_code=500,
        content={"detail": f"Erreur interne du serveur : {str(exc)}"}
    )

assets_path = os.path.join(os.getcwd(), "frontend", "assets")
avatars_path = os.path.join(assets_path, "avatars")
os.makedirs(assets_path, exist_ok=True)
os.makedirs(avatars_path, exist_ok=True)
app.mount("/assets", StaticFiles(directory=assets_path), name="assets")

NO_CACHE_HEADERS = {
    "Cache-Control": "no-store, no-cache, must-revalidate, max-age=0",
    "Pragma": "no-cache",
    "Expires": "0"
}

def extract_attached_files(files_list) -> tuple[list, list]:
    """Extrait le contenu texte, PDF, Word, Excel, Macros VBA et Images des fichiers joints."""
    extracted_files_data = []
    files_names = []
    if not files_list:
        return extracted_files_data, files_names

    for f in files_list:
        fname = f.name if hasattr(f, "name") else f.get("name", "fichier")
        content = f.content if hasattr(f, "content") else f.get("content", "")
        files_names.append(fname)
        ext = fname.split('.')[-1].lower() if '.' in fname else ''

        if content.startswith("data:"):
            if ext in ['png', 'jpg', 'jpeg', 'webp', 'gif'] or content.startswith("data:image/"):
                extracted_files_data.append({"name": fname, "content": content, "is_image": True})
                continue

            try:
                header, b64data = content.split(",", 1)
                file_bytes = base64.b64decode(b64data)
                extracted_text = ""

                if ext == 'pdf':
                    try:
                        import PyPDF2
                        reader = PyPDF2.PdfReader(io.BytesIO(file_bytes))
                        extracted_text = "\n".join([page.extract_text() for page in reader.pages if page.extract_text()])
                    except Exception as e:
                        extracted_text = f"[Erreur lecture PDF: {str(e)}]"
                elif ext in ['docx', 'odt']:
                    try:
                        import zipfile as zf_doc
                        import xml.etree.ElementTree as ET
                        with zf_doc.ZipFile(io.BytesIO(file_bytes)) as dz:
                            xml_file = "word/document.xml" if ext == "docx" else "content.xml"
                            xml_content = dz.read(xml_file)
                            tree = ET.fromstring(xml_content)
                            extracted_text = "".join(tree.itertext())
                    except Exception as e:
                        extracted_text = f"[Erreur lecture document {ext}: {str(e)}]"
                elif ext in ['xls', 'xlsx', 'xlsm', 'xlsb']:
                    try:
                        import openpyxl
                        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=False)
                        for sheet_name in wb.sheetnames:
                            sheet = wb[sheet_name]
                            extracted_text += f"\n--- Feuille : {sheet_name} ---\n"
                            for row in sheet.iter_rows(values_only=True):
                                row_vals = [str(cell) if cell is not None else "" for cell in row]
                                if any(row_vals):
                                    extracted_text += "\t".join(row_vals) + "\n"
                        if ext in ['xlsm', 'xlsb', 'xls']:
                            try:
                                from oletools.olevba import VBA_Parser
                                vbaparser = VBA_Parser("filename", data=file_bytes)
                                if vbaparser.detect_vba_macros():
                                    extracted_text += "\n\n--- MACROS VBA DETECTEES ---\n"
                                    for (_, _, vba_filename, vba_code) in vbaparser.extract_macros():
                                        extracted_text += f"\n// Module: {vba_filename}\n{vba_code}\n"
                            except Exception:
                                pass
                    except Exception as e:
                        extracted_text = f"[Erreur lecture Excel: {str(e)}]"
                else:
                    extracted_text = file_bytes.decode('utf-8', errors='replace')
                
                extracted_files_data.append({"name": fname, "content": extracted_text, "is_image": False})
            except Exception as e:
                extracted_files_data.append({"name": fname, "content": f"[ERREUR DE DECODAGE: {str(e)}]", "is_image": False})
        else:
            extracted_files_data.append({"name": fname, "content": content, "is_image": False})

    return extracted_files_data, files_names

async def scheduler_task():
    tz = pytz.timezone('Europe/Paris')
    while True:
        now = datetime.now(tz)
        if (now.hour == 0 or now.hour == 12) and now.minute == 0:
            db = SessionLocal()
            try:
                settings = db.query(SystemSettings).first()
                if settings:
                    try: await sync_providers_models(db, settings, "Automatique")
                    except: pass
                
                cutoff = datetime.now(timezone.utc) - timedelta(days=7)
                old_projects = db.query(Project).filter(Project.is_pinned == False, Project.created_at < cutoff).all()
                if old_projects:
                    log_activity(f"[Nettoyage BDD] Suppression de {len(old_projects)} discussion(s) > 7 jours.")
                    for op in old_projects:
                        db.delete(op)
                    db.commit()

                cleanup_old_logs(30)
            except Exception as e:
                db.rollback()
                log_activity(f"[Erreur Nettoyage] {str(e)}")
            finally:
                db.close()
            
            await asyncio.sleep(60)
        await asyncio.sleep(30)

@app.on_event("startup")
async def startup_event():
    cleanup_old_logs(30)
    db = SessionLocal()
    try:
        purge_unauthorized_models(db)
    finally:
        db.close()
    asyncio.create_task(scheduler_task())

def is_setup_required(db: Session) -> bool: return db.query(User).filter(User.is_superadmin == True).first() is None

def get_current_user(request: Request, db: Session = Depends(get_db)):
    token = request.cookies.get("session_token")
    if not token: raise HTTPException(status_code=401, detail="Non authentifié")
    payload = verify_token(token)
    if not payload: raise HTTPException(status_code=401, detail="Session expirée")
    user = db.query(User).filter(User.username == payload.get("sub")).first()
    if not user: raise HTTPException(status_code=401, detail="Utilisateur introuvable")
    return user

@app.get("/")
def read_root(db: Session = Depends(get_db)): return RedirectResponse(url="/setup") if is_setup_required(db) else RedirectResponse(url="/login")
@app.get("/setup")
def setup_page(db: Session = Depends(get_db)): return RedirectResponse(url="/login") if not is_setup_required(db) else FileResponse(os.path.join(os.getcwd(), "frontend", "index.html"), headers=NO_CACHE_HEADERS)
@app.get("/login")
def login_page(db: Session = Depends(get_db)): return RedirectResponse(url="/setup") if is_setup_required(db) else FileResponse(os.path.join(os.getcwd(), "frontend", "login.html"), headers=NO_CACHE_HEADERS)

@app.post("/api/setup")
def create_admin(admin_data: AdminCreate, db: Session = Depends(get_db)):
    if not is_setup_required(db): raise HTTPException(status_code=403, detail="Déjà installé.")
    new_admin = User(email=admin_data.email, username=admin_data.username, hashed_password=get_password_hash(admin_data.password), totp_secret=generate_totp_secret(), is_admin=True, is_superadmin=True)
    db.add(new_admin)
    new_settings = SystemSettings(smtp_host=admin_data.smtp_host, smtp_port=admin_data.smtp_port, smtp_user=admin_data.smtp_user, smtp_password=admin_data.smtp_password, openrouter_api_key=admin_data.openrouter_api_key, openrouter_management_key=admin_data.openrouter_management_key, groq_api_key=admin_data.groq_api_key, gemini_api_key=admin_data.gemini_api_key, deepseek_api_key=admin_data.deepseek_api_key, mistral_api_key=admin_data.mistral_api_key, cloudflare_account_id=admin_data.cloudflare_account_id, cloudflare_api_token=admin_data.cloudflare_api_token, huggingface_api_key=admin_data.huggingface_api_key)
    db.add(new_settings)
    db.commit()
    db.refresh(new_admin)
    return {"message": "Succès", "totp_secret": new_admin.totp_secret, "totp_uri": get_totp_uri(new_admin.totp_secret, new_admin.username)}

@app.post("/api/login")
def login(login_data: LoginRequest, response: Response, db: Session = Depends(get_db)):
    user = db.query(User).filter(User.username == login_data.username).first()
    if not user or not verify_password(login_data.password, user.hashed_password): raise HTTPException(status_code=401, detail="Identifiants incorrects.")
    if not verify_totp(user.totp_secret, login_data.totp_code): raise HTTPException(status_code=401, detail="2FA invalide.")
    # Cookie de session fixé à 10 heures (36000 secondes)
    response.set_cookie(key="session_token", value=create_access_token(data={"sub": user.username}), httponly=True, max_age=36000, samesite="lax")
    return {"message": "Connexion réussie"}

@app.get("/dashboard")
def dashboard(request: Request):
    token = request.cookies.get("session_token")
    if not token or not verify_token(token): return RedirectResponse(url="/login")
    return FileResponse(os.path.join(os.getcwd(), "frontend", "dashboard.html"), headers=NO_CACHE_HEADERS)

# --- INTÉGRATION HERMES AGENT (SYNCHRONISÉE AVEC LES PROJETS AETHAS38) ---
@app.get("/api/hermes/status")
async def check_hermes_status(current_user: User = Depends(get_current_user)):
    """Vérifie la disponibilité de l'instance Hermes auto-hébergée."""
    headers = {"Authorization": f"Bearer {HERMES_API_KEY}"}
    async with httpx.AsyncClient(timeout=5.0) as client:
        for endpoint in ["/health", "/v1/models", "/"]:
            try:
                resp = await client.get(f"{HERMES_API_URL}{endpoint}", headers=headers)
                if resp.status_code in [200, 401, 404]:
                    return {"status": "online", "url": HERMES_API_URL}
            except Exception:
                continue
    return {"status": "offline", "url": HERMES_API_URL}

async def _call_hermes_instance(db: Session, project: Project, history: list, extra_files_context: str, settings: SystemSettings, fallback_worker: str) -> str:
    """Appelle Hermes avec l'historique complet de la discussion AETHAS38 pour conserver la mémoire et s'adapter."""
    hermes_messages = [
        {
            "role": "system",
            "content": (
                f"Tu es Hermes, l'agent autonome intégré à AETHAS38 (Discussion: '{project.title}', ID: {project.id}). "
                "Tu partages exactement le même historique de discussion que le mode Multi-Codage AETHAS38. "
                "Garde en mémoire tout le contexte des échanges précédents et adapte-toi au fur et à mesure."
            )
        }
    ]
    for m in history:
        hermes_messages.append({"role": m.role, "content": m.content or ""})

    if extra_files_context and hermes_messages[-1]["role"] == "user":
        hermes_messages[-1]["content"] += f"\n\n[Contenu extrait des fichiers joints] :\n{extra_files_context}"

    headers = {
        "Authorization": f"Bearer {HERMES_API_KEY}",
        "X-API-Key": HERMES_API_KEY,
        "Content-Type": "application/json"
    }

    async with httpx.AsyncClient(timeout=180.0) as client:
        # 1. Endpoint compatible OpenAI v1 (/v1/chat/completions) d'Hermes Gateway
        try:
            payload_v1 = {
                "model": "hermes-agent",
                "messages": hermes_messages,
                "user": f"aethas38_project_{project.id}"
            }
            resp = await client.post(f"{HERMES_API_URL}/v1/chat/completions", headers=headers, json=payload_v1)
            if resp.status_code == 200:
                data = resp.json()
                choices = data.get("choices") or []
                if choices and choices[0].get("message", {}).get("content"):
                    log_activity(f"[Hermes Agent] Réponse synchronisée reçue sur le projet #{project.id}.")
                    return choices[0]["message"]["content"]
        except asyncio.CancelledError:
            raise
        except Exception as e:
            log_activity(f"[Hermes v1] Repli endpoint ({str(e)[:60]})...")

        # 2. Endpoint alternatif (/api/chat) d'Hermes
        try:
            payload_alt = {
                "prompt": hermes_messages[-1]["content"],
                "session_id": f"aethas38_project_{project.id}",
                "history": hermes_messages[:-1]
            }
            resp2 = await client.post(f"{HERMES_API_URL}/api/chat", headers=headers, json=payload_alt)
            if resp2.status_code == 200:
                data2 = resp2.json()
                ans = data2.get("response") or data2.get("message") or data2.get("content")
                if ans:
                    log_activity(f"[Hermes Agent] Réponse reçue via /api/chat sur le projet #{project.id}.")
                    return str(ans)
        except asyncio.CancelledError:
            raise
        except Exception as e:
            log_activity(f"⚠️ [Hermes] Instance locale injoignable ou quota dépassé ({str(e)[:70]}). Bascule sur IA équivalente AETHAS38...")

    # 3. Bascule automatique équivalente si Hermes ou son modèle sous-jacent dépasse son quota
    fallback_ans = await ask_agent_with_fallback(
        db=db,
        settings=settings,
        model_id=fallback_worker or "Qwen/Qwen2.5-Coder-32B-Instruct",
        messages=hermes_messages,
        is_worker=True,
        enable_web=True,
        target_domain="Code & Logique"
    )
    return f"*(Bascule auto AETHAS38 - Quota/Relais Hermes)*\n\n{fallback_ans}"

@app.post("/api/projects/{project_id}/hermes", response_model=List[MessageResponse])
async def create_hermes_project_message(project_id: int, message: MessageCreate, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    """Enregistre et synchronise la discussion Hermes dans le même projet que AETHAS38 Multi-Codage."""
    p = db.query(Project).filter(Project.id == project_id, Project.user_id == current_user.id).first()
    if not p:
        raise HTTPException(status_code=404, detail="Projet introuvable.")

    if message.config:
        p.config_json = json.dumps(message.config.dict())
        db.commit()

    extracted_files_data, files_names = extract_attached_files(message.files)
    extra_files_context = "\n\n".join([f"--- {fd['name']} ---\n{fd['content']}" for fd in extracted_files_data if not fd.get("is_image")])

    db_content = message.content
    if files_names:
        db_content += f"\n\n[Fichiers joints pour analyse : {', '.join(files_names)}]"

    db.add(Message(role="user", content=db_content, project_id=project_id))
    db.commit()

    if p.title == "Nouvelle discussion":
        first_content = message.content.split("\n\n[Fichiers joints")[0].strip()
        new_title = first_content.split('\n')[0][:35].strip() or "Discussion Hermes"
        p.title = new_title + ("..." if len(first_content) > 35 else "")
        db.commit()

    history = db.query(Message).filter(Message.project_id == project_id).order_by(Message.created_at.asc()).all()
    settings = db.query(SystemSettings).first()
    fallback_worker = (message.config.workers[0] if message.config and message.config.workers else "Qwen/Qwen2.5-Coder-32B-Instruct")

    task = asyncio.current_task()
    active_project_tasks[project_id] = task
    try:
        hermes_resp = await _call_hermes_instance(db, p, history, extra_files_context, settings, fallback_worker)
        db.add(Message(role="assistant", content=hermes_resp, project_id=project_id))
        db.commit()
    except asyncio.CancelledError:
        db.add(Message(role="assistant", content="🛑 *[Requête Hermes interrompue par l'utilisateur]*", project_id=project_id))
        db.commit()
    finally:
        active_project_tasks.pop(project_id, None)

    return db.query(Message).filter(Message.project_id == project_id).order_by(Message.created_at.asc()).all()

@app.post("/api/projects/{project_id}/abort")
async def abort_project_request(project_id: int, current_user: User = Depends(get_current_user)):
    """Interrompt immédiatement la requête IA ou Hermes en cours sur le projet."""
    task = active_project_tasks.get(project_id)
    if task and not task.done():
        task.cancel()
        log_activity(f"🛑 Signal d'interruption envoyé sur le projet #{project_id} par {current_user.username}.")
        return {"status": "aborted", "message": "Requête interrompue avec succès."}
    return {"status": "idle", "message": "Aucune requête active à interrompre."}

@app.get("/api/logs")
def get_logs(current_user: User = Depends(get_current_user)):
    return {"logs": activity_logs}

@app.post("/api/logs")
def add_frontend_log(req: LogRequest, current_user: User = Depends(get_current_user)):
    log_activity(f"[Système UI] {req.message}")
    return {"status": "ok"}

@app.get("/api/logs/download")
def download_current_logs(current_user: User = Depends(get_current_user)):
    filename = f"aethas38_logs_{datetime.now().strftime('%Y%m%d_%H%M%S')}.log"
    if os.path.exists(log_file):
        return FileResponse(path=log_file, media_type="text/plain;charset=utf-8", filename=filename, headers=NO_CACHE_HEADERS)
    content = "\n".join(activity_logs) if activity_logs else "Aucun log enregistré."
    return Response(content=content, media_type="text/plain;charset=utf-8", headers={"Content-Disposition": f"attachment; filename={filename}", **NO_CACHE_HEADERS})

@app.post("/api/media/generate", response_model=MediaGenerateResponse)
async def api_generate_media(req: MediaGenerateRequest, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    settings = db.query(SystemSettings).first()
    if not settings:
        raise HTTPException(status_code=400, detail="Paramètres système non initialisés.")
    return await generate_media(
        db=db,
        settings=settings,
        prompt=req.prompt,
        media_type=req.media_type,
        model_id=req.model_id,
        prompter_id=req.prompter_id or "gemini-3.5-flash-lite",
        optimize_prompt=req.optimize_prompt if req.optimize_prompt is not None else True
    )

@app.post("/api/models/suggest_replacement")
async def suggest_replacement(req: ModelReplacementRequest, db: Session = Depends(get_db)):
    log_activity(f"⚠️ Modèle indisponible ou hors quota: {req.missing_model}. Recherche d'un équivalent strict...")
    fallback = find_quota_fallback_model(db, req.missing_model, {req.missing_model}, is_worker=True)
    if fallback:
        return {"suggestion": fallback, "reason": "Remplacement automatique respectant strictement les pré-requis (Gratuit/Payant, Domaine et Contexte)."}
    return {"suggestion": "Qwen/Qwen2.5-Coder-32B-Instruct", "reason": "Repli par défaut."}

@app.get("/api/users/me")
def get_me(current_user: User = Depends(get_current_user)):
    return {"username": current_user.username, "is_admin": current_user.is_admin, "is_superadmin": current_user.is_superadmin, "avatar_path": current_user.avatar_path}

@app.put("/api/users/me/password")
def change_password(passwords: PasswordChange, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    if not verify_password(passwords.old_password, current_user.hashed_password): raise HTTPException(status_code=400, detail="L'ancien mot de passe est incorrect.")
    if len(passwords.new_password) < 8: raise HTTPException(status_code=400, detail="8 caractères minimum.")
    current_user.hashed_password = get_password_hash(passwords.new_password)
    db.commit()
    return {"message": "Mot de passe mis à jour."}

@app.post("/api/users/me/avatar")
async def upload_avatar(file: UploadFile = File(...), db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    file_location = os.path.join(avatars_path, f"user_{current_user.id}.jpg")
    with open(file_location, "wb") as buffer: shutil.copyfileobj(file.file, buffer)
    current_user.avatar_path = f"/assets/avatars/user_{current_user.id}.jpg?v={int(datetime.now().timestamp())}"
    db.commit()
    return {"message": "Avatar mis à jour", "avatar_path": current_user.avatar_path}

@app.get("/api/settings", response_model=SystemSettingsResponse)
def get_settings(db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    if not current_user.is_superadmin: raise HTTPException(status_code=403, detail="Super-Admin requis.")
    return db.query(SystemSettings).first()

@app.put("/api/settings")
def update_settings(settings_data: SystemSettingsUpdate, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    if not current_user.is_superadmin: raise HTTPException(status_code=403, detail="Super-Admin requis.")
    s = db.query(SystemSettings).first()
    if not s: raise HTTPException(status_code=404)
    for k, v in settings_data.dict(exclude_unset=True).items():
        setattr(s, k, v)
    db.commit()
    log_activity("Configuration système mise à jour par le Super-Admin.")
    return {"message": "Paramètres mis à jour avec succès."}

@app.get("/api/projects", response_model=List[ProjectResponse])
def get_projects(db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    return db.query(Project).filter(Project.user_id == current_user.id).order_by(Project.created_at.desc()).all()

@app.post("/api/projects", response_model=ProjectResponse)
def create_project(project: ProjectCreate, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    p = Project(title=project.title, user_id=current_user.id)
    db.add(p); db.commit(); db.refresh(p)
    return p

@app.put("/api/projects/{project_id}/config", response_model=ProjectResponse)
def update_project_config(project_id: int, cfg: ProjectConfigUpdate, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    p = db.query(Project).filter(Project.id == project_id, Project.user_id == current_user.id).first()
    if not p: raise HTTPException(status_code=404)
    p.config_json = cfg.config_json
    db.commit(); db.refresh(p)
    return p

@app.put("/api/projects/{project_id}/rename", response_model=ProjectResponse)
def rename_project(project_id: int, project_data: ProjectRename, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    p = db.query(Project).filter(Project.id == project_id, Project.user_id == current_user.id).first()
    if not p: raise HTTPException(status_code=404)
    p.title = project_data.title; db.commit(); db.refresh(p); return p

@app.put("/api/projects/{project_id}/pin", response_model=ProjectResponse)
def toggle_pin_project(project_id: int, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    p = db.query(Project).filter(Project.id == project_id, Project.user_id == current_user.id).first()
    if not p: raise HTTPException(status_code=404)
    p.is_pinned = not p.is_pinned; db.commit(); db.refresh(p); return p

@app.delete("/api/projects/{project_id}")
def delete_project(project_id: int, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    p = db.query(Project).filter(Project.id == project_id, Project.user_id == current_user.id).first()
    if not p: raise HTTPException(status_code=404)
    db.delete(p); db.commit(); return {"message": "Supprimé"}

@app.get("/api/projects/{project_id}/export")
def export_project(project_id: int, format: str = "txt", db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    p = db.query(Project).filter(Project.id == project_id, Project.user_id == current_user.id).first()
    messages = db.query(Message).filter(Message.project_id == project_id).order_by(Message.created_at.asc()).all()
    if format == "json":
        data = [{"role": m.role, "content": m.content, "date": m.created_at.isoformat()} for m in messages]
        return Response(content=json.dumps(data, indent=2), media_type="application/json", headers={"Content-Disposition": f"attachment; filename=export_{project_id}.json"})
    text_out = f"--- HISTORIQUE : {p.title} ---\n\n"
    for m in messages: text_out += f"[{m.created_at.strftime('%Y-%m-%d %H:%M:%S')}] {'VOUS' if m.role == 'user' else 'AETHAS38 / HERMES'}:\n{m.content}\n\n{'-'*50}\n\n"
    return Response(content=text_out, media_type="text/plain;charset=utf-8", headers={"Content-Disposition": f"attachment; filename=export_{project_id}.txt"})

@app.get("/api/projects/{project_id}/messages", response_model=List[MessageResponse])
def get_messages(project_id: int, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    return db.query(Message).filter(Message.project_id == project_id).order_by(Message.created_at.asc()).all()

@app.post("/api/projects/{project_id}/messages", response_model=List[MessageResponse])
async def create_message(project_id: int, message: MessageCreate, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    p = db.query(Project).filter(Project.id == project_id, Project.user_id == current_user.id).first()
    if not p: raise HTTPException(status_code=404, detail="Projet introuvable.")

    if message.config:
        p.config_json = json.dumps(message.config.dict())
        db.commit()

    extracted_files_data, files_names = extract_attached_files(message.files)

    db_content = message.content
    if files_names:
        db_content += f"\n\n[Fichiers joints pour analyse : {', '.join(files_names)}]"

    db.add(Message(role=message.role, content=db_content, project_id=project_id))
    db.commit()

    if p.title == "Nouvelle discussion":
        user_msgs = db.query(Message).filter(Message.project_id == project_id, Message.role == "user").order_by(Message.created_at.asc()).all()
        if len(user_msgs) >= 1:
            first_content = user_msgs[0].content.split("\n\n[Fichiers joints")[0].strip()
            new_title = first_content.split('\n')[0][:35].strip() or "Discussion"
            p.title = new_title + ("..." if len(first_content) > 35 else "")
            db.commit()
    
    history = db.query(Message).filter(Message.project_id == project_id).order_by(Message.created_at.asc()).all()
    settings = db.query(SystemSettings).first()
    conf = message.config.dict() if message.config else {"workers": ["Qwen/Qwen2.5-Coder-32B-Instruct"]} 
    
    task = asyncio.current_task()
    active_project_tasks[project_id] = task
    try:
        ai_resp = await run_orchestrator(db, history, settings, conf, extracted_files_data)
        db.add(Message(role="assistant", content=ai_resp, project_id=project_id))
        db.commit()
    except asyncio.CancelledError:
        db.add(Message(role="assistant", content="🛑 *[Requête interrompue par l'utilisateur]*", project_id=project_id))
        db.commit()
    finally:
        active_project_tasks.pop(project_id, None)

    return db.query(Message).filter(Message.project_id == project_id).order_by(Message.created_at.asc()).all()

@app.get("/api/models/info")
def get_models_info(db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    settings = db.query(SystemSettings).first()
    return {
        "last_sync_date": settings.last_sync_date.isoformat() if settings and settings.last_sync_date else None,
        "last_sync_type": settings.last_sync_type if settings else None,
        "models": db.query(AIModel).order_by(AIModel.name.asc()).all(),
        "finances": db.query(FinancialLog).all()
    }

@app.get("/api/finances")
def get_finances(db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    return db.query(FinancialLog).all()

@app.post("/api/models/sync")
async def trigger_model_sync(db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    if not current_user.is_admin: raise HTTPException(status_code=403, detail="Accès admin requis.")
    settings = db.query(SystemSettings).first()
    return await sync_providers_models(db, settings, "Manuelle")

@app.get("/api/models/export")
def export_models(db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    if not current_user.is_admin: raise HTTPException(status_code=403, detail="Accès admin requis.")
    models = db.query(AIModel).order_by(AIModel.provider.asc(), AIModel.name.asc()).all()

    csv_io = io.StringIO()
    writer = csv.writer(csv_io, delimiter=',')
    writer.writerow(["Provider", "Model ID", "Name", "Domain", "Is Free", "Context Length", "Pricing Prompt", "Pricing Completion", "Description"])
    for m in models:
        writer.writerow([m.provider, m.model_id, m.name, m.domain, m.is_free, m.context_length, m.pricing_prompt, m.pricing_completion, m.description_fr])
    
    md_content = f"# Extraction des Modèles IA - AETHAS38\n\n**Date d'extraction :** {datetime.now().strftime('%d/%m/%Y à %H:%M:%S')}\n\n"
    providers = sorted(list(set(m.provider for m in models)))
    for prov in providers:
        md_content += f"## Fournisseur : {prov.upper()}\n\n"
        prov_models = [m for m in models if m.provider == prov]
        for m in prov_models:
            price_info = "**GRATUIT**" if m.is_free else f"In: ${m.pricing_prompt:.2f} / Out:${m.pricing_completion:.2f}"
            ctx_info = f"{int((m.context_length or 0)/1000)}k"
            desc = m.description_fr.replace('\n', ' ') if m.description_fr else ""
            md_content += f"- **{m.name or m.model_id}** (`{m.model_id}`)\n"
            md_content += f"  - *Domaine :* {m.domain}\n"
            md_content += f"  - *Prix (1M tokens) :* {price_info}\n"
            md_content += f"  - *Contexte :* {ctx_info}\n"
            md_content += f"  - *Description :* {desc}\n\n"

    zip_io = io.BytesIO()
    with zipfile.ZipFile(zip_io, mode='w', compression=zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("models_export.csv", csv_io.getvalue().encode('utf-8'))
        zf.writestr(f"{datetime.now().strftime('%Y%m%d')}-extraction-modeles.md", md_content.encode('utf-8'))

    zip_io.seek(0)
    return Response(
        content=zip_io.getvalue(), 
        media_type="application/zip", 
        headers={"Content-Disposition": f"attachment; filename=aethas38_models_{datetime.now().strftime('%Y%m%d')}.zip"}
    )